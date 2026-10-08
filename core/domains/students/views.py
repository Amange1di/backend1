from io import BytesIO

from django.db import models, transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework import status, viewsets
from rest_framework.authtoken.models import Token
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied
from rest_framework.response import Response
from django.http import HttpResponse

from core.models import Group, Student, User
from core.audit import write_audit
from core.permissions import (
    IsCourseAdminOrManagerOrStudentReadOnly,
)

from .serializers import (
    StudentSerializer,
    TransferGroupSerializer,
)
from .services import sync_student_user
from core.domains.auth.first_login import issue_first_login_password, create_first_login_link_token


class StudentViewSet(viewsets.ModelViewSet):
    queryset = (
        Student.objects.filter(
            archived_at__isnull=True
        )
        .select_related(
            "user",
            "company",
            "primary_course",
        )
        .prefetch_related("groups")
        .order_by("-created_at")
    )
    serializer_class = StudentSerializer
    permission_classes = [
        IsCourseAdminOrManagerOrStudentReadOnly
    ]

    def get_queryset(self):
        queryset = super().get_queryset()
        user = self.request.user
        selected_branch = self.request.COOKIES.get("eduosh_branch")
        branch_id = int(selected_branch) if selected_branch and selected_branch.isdigit() else None
        if user.is_authenticated and user.role in (User.Role.COURSE_ADMIN, User.Role.MANAGER):
            allowed = user.branches.filter(is_active=True)
            if branch_id:
                if not allowed.filter(id=branch_id).exists():
                    raise PermissionDenied("branch_access_denied")
                queryset = queryset.filter(groups__branch_id=branch_id).distinct()
            else:
                queryset = queryset.filter(groups__branch__in=allowed).distinct()
        elif branch_id and user.is_authenticated and user.role == User.Role.COMPANY_OWNER:
            queryset = queryset.filter(groups__branch_id=branch_id).distinct()

        group_id = self.request.query_params.get("group")
        if group_id:
            try:
                queryset = queryset.filter(groups__id=int(group_id)).distinct()
            except (TypeError, ValueError):
                return queryset.none()

        if (
            user.is_authenticated
            and user.role in (User.Role.COMPANY_OWNER, User.Role.COURSE_ADMIN)
        ):
            return queryset.filter(
                models.Q(company=user.company)
                | models.Q(
                    primary_course__admins=user
                )
                | models.Q(
                    groups__course__admins=user
                )
            ).distinct()

        if (
            user.is_authenticated
            and user.role == User.Role.MANAGER
        ):
            if not user.company:
                return queryset.none()

            return queryset.filter(
                models.Q(company=user.company)
                | models.Q(
                    primary_course__admins=user
                )
                | models.Q(
                    groups__company=user.company
                )
                | models.Q(
                    groups__course__admins=user
                )
            ).distinct()

        if (
            user.is_authenticated
            and user.role == User.Role.TEACHER
        ):
            return queryset.filter(
                groups__teacher=user
            ).distinct()

        if (
            user.is_authenticated
            and user.role == User.Role.STUDENT
        ):
            return queryset.filter(user=user)

        return queryset.none()

    def _validate_course_access(
        self,
        *,
        user,
        course,
    ):
        if not course:
            return

        allowed = course.admins.filter(
            id=user.id
        ).exists()

        if user.role == User.Role.MANAGER:
            allowed = course.admins.filter(
                company=user.company
            ).exists()

        if not allowed:
            raise PermissionDenied(
                "Not allowed for this course."
            )

    def _validate_group_access(
        self,
        *,
        user,
        groups,
    ):
        for group in groups:
            if group.course:
                allowed = group.course.admins.filter(
                    id=user.id
                ).exists()

                if user.role == User.Role.MANAGER:
                    allowed = (
                        group.course.admins.filter(
                            company=user.company
                        ).exists()
                    )

                if not allowed:
                    raise PermissionDenied(
                        "group_access_denied"
                    )

            elif (
                group.company
                and group.company != user.company
            ):
                raise PermissionDenied(
                    "group_access_denied"
                )

    @staticmethod
    def _auto_course_from_groups(
        course,
        groups,
    ):
        if course or not groups:
            return None

        first_course = groups[0].course
        if (
            first_course
            and all(
                group.course_id
                == first_course.id
                for group in groups
            )
        ):
            return first_course

        return None

    def perform_create(self, serializer):
        user = self.request.user

        if user.role in (
            User.Role.COMPANY_OWNER,
            User.Role.COURSE_ADMIN,
            User.Role.MANAGER,
        ):
            course = serializer.validated_data.get(
                "primary_course"
            )
            groups = serializer.validated_data.get(
                "group_ids",
                [],
            )
            account = serializer.validated_data.get(
                "user"
            )

            self._validate_course_access(
                user=user,
                course=course,
            )
            self._validate_group_access(
                user=user,
                groups=groups,
            )

            auto_course = (
                self._auto_course_from_groups(
                    course,
                    groups,
                )
            )

            if (
                account
                and account.company
                and account.company
                != user.company
            ):
                raise PermissionDenied(
                    "Not allowed for this user."
                )

            save_kwargs = {
                "company": user.company
            }

            if course:
                save_kwargs[
                    "primary_course"
                ] = course
            elif auto_course:
                save_kwargs[
                    "primary_course"
                ] = auto_course

            student = serializer.save(
                **save_kwargs
            )
            sync_student_user(
                student,
                created_by=user,
            )
            return

        student = serializer.save()
        sync_student_user(
            student,
            created_by=(
                user
                if user.is_authenticated
                else None
            ),
        )

    def perform_update(self, serializer):
        user = self.request.user

        if user.role in (
            User.Role.COMPANY_OWNER,
            User.Role.COURSE_ADMIN,
            User.Role.MANAGER,
        ):
            if (
                user.role == User.Role.MANAGER
                and "can_login"
                in serializer.validated_data
            ):
                raise PermissionDenied(
                    (
                        "Managers cannot change "
                        "student login access."
                    )
                )

            course = serializer.validated_data.get(
                "primary_course",
                None,
            )
            groups = serializer.validated_data.get(
                "group_ids",
                [],
            )

            self._validate_course_access(
                user=user,
                course=course,
            )
            self._validate_group_access(
                user=user,
                groups=groups,
            )

            auto_course = (
                self._auto_course_from_groups(
                    course,
                    groups,
                )
            )

            save_kwargs = {}
            if course:
                save_kwargs[
                    "primary_course"
                ] = course
            elif auto_course:
                save_kwargs[
                    "primary_course"
                ] = auto_course

            student = serializer.save(
                **save_kwargs
            )
            sync_student_user(
                student,
                created_by=user,
            )
            return

        student = serializer.save()
        sync_student_user(
            student,
            created_by=(
                user
                if user.is_authenticated
                else None
            ),
        )

    def destroy(self, request, *args, **kwargs):
        if request.user.role in (
            User.Role.MANAGER,
            User.Role.STUDENT,
        ):
            raise PermissionDenied(
                "Not allowed to archive students."
            )

        student = self.get_object()
        student.archived_at = timezone.now()
        student.can_login = False
        student.save(
            update_fields=[
                "archived_at",
                "can_login",
            ]
        )

        if student.user_id:
            student.user.is_active = False
            student.user.save(
                update_fields=["is_active"]
            )
            Token.objects.filter(
                user=student.user
            ).delete()

        write_audit(
            request,
            action="student.archived",
            obj=student,
            company=student.company,
            after={
                "archived_at": (
                    student.archived_at.isoformat()
                ),
                "can_login": False,
            },
        )

        return Response(status=204)

    @action(
        detail=False,
        methods=["get"],
        url_path="import-template",
    )
    def import_template(self, request):
        if request.user.role not in (
            User.Role.COURSE_ADMIN,
            User.Role.MANAGER,
        ):
            raise PermissionDenied(
                "staff_only"
            )

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Students"

        headers = [
            "first_name",
            "last_name",
            "phone",
            "telegram",
            "notes",
        ]
        sheet.append(headers)
        sheet.append(
            [
                "Аман",
                "Майрамбек уулу",
                "+996700000000",
                "@aman",
                "Комментарий",
            ]
        )

        sheet.freeze_panes = "A2"
        widths = {
            "A": 22,
            "B": 24,
            "C": 20,
            "D": 20,
            "E": 36,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )
        response[
            "Content-Disposition"
        ] = 'attachment; filename="students_import_template.xlsx"'
        return response

    @action(
        detail=False,
        methods=["post"],
        url_path="bulk-import",
    )
    def bulk_import(self, request):
        user = request.user
        if user.role not in (
            User.Role.COURSE_ADMIN,
            User.Role.MANAGER,
        ):
            raise PermissionDenied(
                "staff_only"
            )

        if not user.company:
            return Response(
                {"detail": "company_not_found"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        upload = request.FILES.get("file")
        if not upload:
            return Response(
                {"detail": "excel_file_required"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if not upload.name.lower().endswith(".xlsx"):
            return Response(
                {"detail": "excel_xlsx_only"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        group = None
        group_id = request.data.get("group_id")
        if group_id:
            try:
                group = Group.objects.get(
                    id=group_id,
                    archived_at__isnull=True,
                    company=user.company,
                )
            except (Group.DoesNotExist, ValueError, TypeError):
                return Response(
                    {"detail": "group_not_found"},
                    status=status.HTTP_400_BAD_REQUEST,
                )

        try:
            workbook = load_workbook(
                upload,
                read_only=True,
                data_only=True,
            )
        except Exception:
            return Response(
                {"detail": "excel_read_failed"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return Response(
                {"detail": "excel_empty"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        headers = [
            str(value).strip() if value is not None else ""
            for value in rows[0]
        ]
        expected = {
            "first_name",
            "last_name",
            "phone",
            "telegram",
            "notes",
        }
        missing = sorted(expected - set(headers))
        if missing:
            return Response(
                {
                    "detail": "excel_missing_columns",
                    "missing_columns": missing
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        if len(rows) - 1 > 1000:
            return Response(
                {"detail": "student_import_limit_exceeded", "max": 1000},
                status=status.HTTP_400_BAD_REQUEST,
            )

        index = {name: headers.index(name) for name in expected}
        parsed_rows = []
        errors = []
        phones_seen = set()
        existing_phones = set(
            Student.objects.filter(
                company=user.company,
                archived_at__isnull=True,
            ).values_list("phone", flat=True)
        )

        for excel_row, values in enumerate(rows[1:], start=2):
            if not values or all(value in (None, "") for value in values):
                continue

            first_name = str(
                values[index["first_name"]] or ""
            ).strip()
            last_name = str(
                values[index["last_name"]] or ""
            ).strip()
            phone = str(values[index["phone"]] or "").strip()
            telegram = str(
                values[index["telegram"]] or ""
            ).strip()
            notes = str(values[index["notes"]] or "").strip()

            if not first_name:
                errors.append(
                    {"row": excel_row, "field": "first_name", "error_key": "first_name_required"}
                )
            if not phone:
                errors.append(
                    {"row": excel_row, "field": "phone", "error_key": "phone_required"}
                )
            elif phone in phones_seen:
                errors.append(
                    {
                        "row": excel_row,
                        "field": "phone",
                        "error_key": "duplicate_phone_in_file",
                    }
                )
            elif phone in existing_phones:
                errors.append(
                    {
                        "row": excel_row,
                        "field": "phone",
                        "error_key": "student_phone_exists",
                    }
                )

            phones_seen.add(phone)

            parsed_rows.append(
                {
                    "row": excel_row,
                    "first_name": first_name,
                    "last_name": last_name,
                    "phone": phone,
                    "telegram": telegram,
                    "notes": notes,
                }
            )

        if not parsed_rows:
            return Response(
                {"detail": "student_import_no_rows"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if errors:
            return Response(
                {
                    "detail": "excel_validation_failed",
                    "errors": errors[:100],
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        created_students = []
        with transaction.atomic():
            for item in parsed_rows:
                serializer = StudentSerializer(
                    data={
                        "first_name": item["first_name"],
                        "last_name": item["last_name"],
                        "phone": item["phone"],
                        "telegram": item["telegram"],
                        "notes": item["notes"],
                        "group_ids": [group.id] if group else [],
                        "primary_course": (
                            group.course_id
                            if group and group.course_id
                            else None
                        ),
                    },
                    context={"request": request},
                )
                serializer.is_valid(raise_exception=True)

                if group:
                    self._validate_group_access(
                        user=user,
                        groups=[group],
                    )

                save_kwargs = {"company": user.company}
                if group and group.course:
                    save_kwargs["primary_course"] = group.course

                student = serializer.save(**save_kwargs)
                sync_student_user(
                    student,
                    created_by=user,
                )
                created_students.append(student)

        return Response(
            {
                "created": len(created_students),
                "group": (
                    {
                        "id": group.id,
                        "name": group.name,
                    }
                    if group
                    else None
                ),
            },
            status=status.HTTP_201_CREATED,
        )

    @action(
        detail=True,
        methods=["post"],
        url_path="reset-password",
    )
    def reset_password(
        self,
        request,
        pk=None,
    ):
        if request.user.role not in (
            User.Role.COURSE_ADMIN,
            User.Role.MANAGER,
        ):
            raise PermissionDenied(
                (
                    "staff_only"
                )
            )

        student = self.get_object()

        if (
            student.company
            != request.user.company
        ):
            raise PermissionDenied(
                "student_access_denied"
            )

        if not student.user:
            sync_student_user(
                student,
                created_by=request.user,
            )
            student.refresh_from_db()

        Token.objects.filter(
            user=student.user
        ).delete()
        one_time_password = (
            issue_first_login_password(
                student.user
            )
        )

        return Response(
            {
                "detail": (
                    "student_password_reset"
                ),
                "must_set_password": True,
                "login": student.phone,
                "username": (
                    student.user.username
                ),
                "one_time_password": (
                    one_time_password
                ),
                "first_login_token": (
                    create_first_login_link_token(
                        student.user
                    )
                ),
            }
        )

    @action(
        detail=True,
        methods=["post"],
        url_path="transfer-group",
    )
    def transfer_group(
        self,
        request,
        pk=None,
    ):
        student = self.get_object()
        serializer = TransferGroupSerializer(
            data=request.data
        )
        serializer.is_valid(
            raise_exception=True
        )

        new_group = serializer.validated_data[
            "new_group"
        ]
        note = serializer.validated_data.get(
            "note",
            "",
        )

        user = request.user
        if user.role not in (
            User.Role.COURSE_ADMIN,
            User.Role.MANAGER,
        ):
            raise PermissionDenied(
                (
                    "staff_only"
                )
            )

        if student.company != new_group.company:
            raise PermissionDenied(
                (
                    "student_group_company_mismatch"
                )
            )

        if student.company != user.company:
            raise PermissionDenied(
                "student_access_denied"
            )

        if new_group.company != user.company:
            raise PermissionDenied(
                "group_access_denied"
            )

        student.groups.clear()
        student.groups.add(new_group)

        if new_group.course:
            student.primary_course = (
                new_group.course
            )

        if note:
            existing = student.notes or ""
            timestamp = timezone.now().strftime(
                "%d.%m.%Y %H:%M"
            )
            transfer_note = (
                f"[{timestamp}] Переведён в группу "
                f"«{new_group.name}». {note}"
            )
            student.notes = (
                f"{transfer_note}\n{existing}"
                if existing
                else transfer_note
            )

        student.save(
            update_fields=[
                "primary_course",
                "notes",
            ]
        )

        return Response(
            {
                "status": "ok",
                "detail": (
                    "student_transferred"
                ),
            }
        )
