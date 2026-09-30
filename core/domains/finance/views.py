from calendar import monthrange
from datetime import date

from django.db.models import Sum
from django.http import HttpResponse
from django.utils import timezone
from rest_framework import viewsets
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework.exceptions import PermissionDenied

from core.models import Expense, Group, GroupMonth, Payment, Student, User
from core.permissions import (
    IsCourseAdminOrManager,
    IsCourseAdminOrManagerReadOnly,
)

from .serializers import (
    ExpenseSerializer,
    GroupMonthSerializer,
)


class ExpenseViewSet(viewsets.ModelViewSet):
    queryset = Expense.objects.none()
    permission_classes = [
        IsCourseAdminOrManager
    ]
    serializer_class = ExpenseSerializer

    def get_queryset(self):
        user = self.request.user

        if user.role in (
            User.Role.COURSE_ADMIN,
            User.Role.MANAGER,
        ):
            if user.company:
                return Expense.objects.filter(
                    company=user.company
                )

        return Expense.objects.none()

    def perform_create(self, serializer):
        company = self.request.user.company
        if not company:
            raise PermissionDenied(
                (
                    "У вас нет компании "
                    "для создания расходов."
                )
            )

        serializer.save(company=company)


class GroupMonthViewSet(viewsets.ModelViewSet):
    queryset = GroupMonth.objects.none()
    permission_classes = [
        IsCourseAdminOrManagerReadOnly
    ]
    serializer_class = GroupMonthSerializer

    def get_queryset(self):
        user = self.request.user

        if user.role in (
            User.Role.COURSE_ADMIN,
            User.Role.MANAGER,
        ):
            if user.company:
                queryset = GroupMonth.objects.filter(
                    group__company=user.company
                )
            else:
                return GroupMonth.objects.none()

        elif user.role == User.Role.TEACHER:
            queryset = GroupMonth.objects.filter(
                group__teacher=user
            )

        else:
            return GroupMonth.objects.none()

        group_id = self.request.query_params.get(
            "group"
        )
        if group_id and group_id.isdigit():
            queryset = queryset.filter(
                group_id=int(group_id)
            )

        teacher_id = self.request.query_params.get(
            "teacher_id"
        )
        if (
            teacher_id
            and teacher_id.isdigit()
        ):
            queryset = queryset.filter(
                group__teacher_id=int(
                    teacher_id
                )
            )

        month_number = (
            self.request.query_params.get(
                "month_number"
            )
        )
        if (
            month_number
            and month_number.isdigit()
        ):
            queryset = queryset.filter(
                month_number=int(
                    month_number
                )
            )

        status_value = (
            self.request.query_params.get(
                "status"
            )
        )
        if status_value:
            queryset = queryset.filter(
                status=status_value
            )

        return (
            queryset
            .select_related(
                "group",
                "group__course",
            )
            .prefetch_related(
                "group__students"
            )
        )

    def list(self, request, *args, **kwargs):
        group_id = request.query_params.get(
            "group"
        )
        if group_id:
            try:
                group = Group.objects.get(
                    id=group_id
                )
                if (
                    group.total_months
                    and group.total_months > 0
                ):
                    existing = set(
                        GroupMonth.objects.filter(
                            group=group
                        ).values_list(
                            "month_number",
                            flat=True,
                        )
                    )

                    for month_number in range(
                        1,
                        group.total_months + 1,
                    ):
                        if (
                            month_number
                            not in existing
                        ):
                            GroupMonth.objects.create(
                                group=group,
                                month_number=month_number,
                                status=(
                                    GroupMonth.Status.PENDING
                                ),
                            )

            except Group.DoesNotExist:
                pass

        return super().list(
            request,
            *args,
            **kwargs,
        )


class FinanceDashboardView(APIView):
    """
    Дашборд финансовой сводки.
    GET /api/finance/dashboard/
    """
    permission_classes = [IsCourseAdminOrManager]

    def get(self, request):
        from django.db.models import Sum
        
        from calendar import monthrange

        user = request.user
        if user.role == User.Role.COURSE_ADMIN:
            company = user.company
        elif user.role == User.Role.MANAGER:
            company = user.company
        else:
            return Response({'detail': 'Нет доступа'}, status=403)

        if not company:
            return Response({'detail': 'Компания не найдена'}, status=404)

        now = timezone.now().date()
        first_of_month = date(now.year, now.month, 1)
        end_of_month = date(now.year, now.month, monthrange(now.year, now.month)[1])
        start_of_year = date(now.year, 1, 1)

        monthly_income = Payment.objects.filter(
            company=company, status='paid',
            paid_at__gte=first_of_month, paid_at__lte=end_of_month
        ).aggregate(total=Sum('amount'))['total'] or 0

        monthly_expenses = Expense.objects.filter(
            company=company,
            date__gte=first_of_month, date__lte=end_of_month
        ).aggregate(total=Sum('amount'))['total'] or 0

        monthly_salaries = GroupMonth.objects.filter(
            group__company=company, teacher_salary__isnull=False,
            completed_at__gte=first_of_month, completed_at__lte=end_of_month
        ).aggregate(total=Sum('teacher_salary'))['total'] or 0

        monthly_total_expenses = float(monthly_expenses) + float(monthly_salaries)

        yearly_income = Payment.objects.filter(
            company=company, status='paid',
            paid_at__gte=start_of_year, paid_at__lte=end_of_month
        ).aggregate(total=Sum('amount'))['total'] or 0

        yearly_expenses = Expense.objects.filter(
            company=company,
            date__gte=start_of_year, date__lte=end_of_month
        ).aggregate(total=Sum('amount'))['total'] or 0

        yearly_salaries = GroupMonth.objects.filter(
            group__company=company, teacher_salary__isnull=False,
            completed_at__gte=start_of_year, completed_at__lte=end_of_month
        ).aggregate(total=Sum('teacher_salary'))['total'] or 0

        yearly_total_expenses = float(yearly_expenses) + float(yearly_salaries)

        total_debt = Payment.objects.filter(
            company=company, status='debt'
        ).aggregate(total=Sum('amount'))['total'] or 0

        students_count = Student.objects.filter(company=company).count()

        from finance.models import Budget
        active_budgets_count = Budget.objects.filter(company=company, is_active=True).count()

        return Response({
            'monthly': {
                'income': float(monthly_income),
                'expenses': monthly_total_expenses,
                'net': float(monthly_income) - monthly_total_expenses,
            },
            'yearly': {
                'income': float(yearly_income),
                'expenses': yearly_total_expenses,
                'net': float(yearly_income) - yearly_total_expenses,
            },
            'total_debt': float(total_debt),
            'students_count': students_count,
            'active_budgets': active_budgets_count,
        })


class FinanceExportView(APIView):
    """
    Экспорт финансовых данных.
    GET /api/finance/export/<format>/
    Поддерживаемые форматы: json, xlsx, csv
    """
    permission_classes = [IsCourseAdminOrManager]

    def get(self, request, export_format: str):
        from django.db.models import Sum
        from calendar import monthrange
        from django.http import HttpResponse

        user = request.user
        if user.role == User.Role.COURSE_ADMIN:
            company = user.company
        elif user.role == User.Role.MANAGER:
            company = user.company
        else:
            return Response({'detail': 'Нет доступа'}, status=403)

        if not company:
            return Response({'detail': 'Компания не найдена'}, status=404)

        now = timezone.now().date()
        first_of_month = date(now.year, now.month, 1)
        end_of_month = date(now.year, now.month, monthrange(now.year, now.month)[1])

        payments = Payment.objects.filter(
            company=company,
            paid_at__gte=first_of_month, paid_at__lte=end_of_month
        ).select_related('student', 'group')

        expenses = Expense.objects.filter(
            company=company,
            date__gte=first_of_month, date__lte=end_of_month
        )

        if export_format == 'json':
            return Response({
                'payments': list(payments.values('id', 'student__first_name', 'student__last_name', 'amount', 'status', 'paid_at')),
                'expenses': list(expenses.values('id', 'description', 'amount', 'category', 'date')),
            })

        elif export_format == 'xlsx':
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

            wb = openpyxl.Workbook()

            # ─── Sheet 1: Payments ───
            ws1 = wb.active
            ws1.title = "Платежи"

            header_font = Font(bold=True, color="FFFFFF", size=11)
            header_fill = PatternFill(start_color="2D3748", end_color="2D3748", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center")
            thin_border = Border(
                left=Side(style="thin", color="E2E8F0"),
                right=Side(style="thin", color="E2E8F0"),
                top=Side(style="thin", color="E2E8F0"),
                bottom=Side(style="thin", color="E2E8F0"),
            )

            payment_headers = ["ID", "Студент", "Сумма", "Статус", "Дата оплаты"]
            for col_idx, header in enumerate(payment_headers, 1):
                cell = ws1.cell(row=1, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border

            for row_idx, payment in enumerate(payments, 2):
                student_name = f"{payment.student.first_name or ''} {payment.student.last_name or ''}".strip()
                ws1.cell(row=row_idx, column=1, value=payment.id).border = thin_border
                ws1.cell(row=row_idx, column=2, value=student_name or "—").border = thin_border
                ws1.cell(row=row_idx, column=3, value=float(payment.amount)).border = thin_border
                ws1.cell(row=row_idx, column=4, value=payment.get_status_display()).border = thin_border
                ws1.cell(row=row_idx, column=5, value=payment.paid_at.strftime("%d.%m.%Y") if payment.paid_at else "—").border = thin_border

            ws1.column_dimensions["A"].width = 8
            ws1.column_dimensions["B"].width = 35
            ws1.column_dimensions["C"].width = 15
            ws1.column_dimensions["D"].width = 15
            ws1.column_dimensions["E"].width = 15

            # ─── Sheet 2: Expenses ───
            ws2 = wb.create_sheet("Расходы")

            expense_headers = ["ID", "Описание", "Сумма", "Категория", "Дата"]
            for col_idx, header in enumerate(expense_headers, 1):
                cell = ws2.cell(row=1, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = PatternFill(start_color="4A5568", end_color="4A5568", fill_type="solid")
                cell.alignment = header_alignment
                cell.border = thin_border

            expense_category_labels = dict(Expense.Category.choices)

            for row_idx, expense in enumerate(expenses, 2):
                ws2.cell(row=row_idx, column=1, value=expense.id).border = thin_border
                ws2.cell(row=row_idx, column=2, value=expense.description).border = thin_border
                ws2.cell(row=row_idx, column=3, value=float(expense.amount)).border = thin_border
                ws2.cell(row=row_idx, column=4, value=expense_category_labels.get(expense.category, expense.category)).border = thin_border
                ws2.cell(row=row_idx, column=5, value=expense.date.strftime("%d.%m.%Y") if expense.date else "—").border = thin_border

            ws2.column_dimensions["A"].width = 8
            ws2.column_dimensions["B"].width = 40
            ws2.column_dimensions["C"].width = 15
            ws2.column_dimensions["D"].width = 25
            ws2.column_dimensions["E"].width = 15

            response = HttpResponse(
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            response["Content-Disposition"] = f'attachment; filename="finance_{now.strftime("%Y-%m")}.xlsx"'
            wb.save(response)
            return response

        elif export_format == 'csv':
            import csv

            response = HttpResponse(content_type="text/csv; charset=utf-8-sig")
            response["Content-Disposition"] = f'attachment; filename="finance_{now.strftime("%Y-%m")}.csv"'

            writer = csv.writer(response)

            writer.writerow(["=== ПЛАТЕЖИ ==="])
            writer.writerow(["ID", "Студент", "Сумма", "Статус", "Дата оплаты"])
            for payment in payments:
                student_name = f"{payment.student.first_name or ''} {payment.student.last_name or ''}".strip()
                writer.writerow([
                    payment.id,
                    student_name or "—",
                    float(payment.amount),
                    payment.get_status_display(),
                    payment.paid_at.strftime("%d.%m.%Y") if payment.paid_at else "—",
                ])

            writer.writerow([])
            writer.writerow(["=== РАСХОДЫ ==="])
            writer.writerow(["ID", "Описание", "Сумма", "Категория", "Дата"])

            expense_category_labels = dict(Expense.Category.choices)
            for expense in expenses:
                writer.writerow([
                    expense.id,
                    expense.description,
                    float(expense.amount),
                    expense_category_labels.get(expense.category, expense.category),
                    expense.date.strftime("%d.%m.%Y") if expense.date else "—",
                ])

            return response

        else:
            return Response({'detail': f'Формат "{export_format}" не поддерживается. Используйте: json, xlsx, csv'}, status=400)
